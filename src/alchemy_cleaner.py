"""Deterministic cleaner for TT2 Alchemy recipe matrices.

Reward keys come from ``default_importance_scores`` in ``streamlit_app.py``,
read with ``ast`` so the Streamlit app is never imported.
"""

from __future__ import annotations

import ast
import csv
import io
import json
import os
import re
import tempfile
import urllib.error
import urllib.request
from collections import Counter
from dataclasses import asdict, dataclass, field
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter


class CleanerError(Exception):
    """Expected failure. The production CSV is left untouched."""


# Historical spreadsheet wording, matched case-insensitively. Targets must be
# current optimiser keys, otherwise the cell stays unsupported.
REWARD_ALIASES = {
    "pet egg": "Eggs",
    "pet eggs": "Eggs",
    "wild card": "Wildcards",
    "wild cards": "Wildcards",
    "skill point": "Skill Points",
    "skill points": "Skill Points",
    "clan scroll": "Clan Scroll",
    "clan scrolls": "Clan Scroll",
    "fortune scroll": "Fortune Scroll",
    "fortune scrolls": "Fortune Scroll",
    "fortune hero weapons": "Hero Weapons",
    "legendary insult": "Legendary Equipment",
    "safety tongs (event set)": "Event Equipment",
    "unique equipment (elder snow cap)": "Event Equipment",
    "unique equipment (the retaliator reforged)": "Event Equipment",
}

# Whole-cell overrides, applied before compound detection. Only user-approved
# values belong here; any other compound reward still fails.
EXACT_VALUES = {
    "avatar + 108 currency": "108 Currency",
}

EQUIPMENT_VARIANTS = [
    (re.compile(r"^event equipment\s*\(.*\)$", re.IGNORECASE), "Event Equipment"),
    (re.compile(r"^legendary equip(?:ment)?\s*\(.*\)$", re.IGNORECASE), "Legendary Equipment"),
]

ARCHIVE_TERMS = ("archive", "archived", "old", "copy", "unsorted")
VALID = ("ingredient", "supported_reward")

QUANTITY_RE = re.compile(r"^x?(\d+)\s+(.+)$", re.IGNORECASE)
CANONICAL_RE = re.compile(r"^([1-9]\d*) (.+)$")
COMPOUND_RE = re.compile(r"\S\s*\+\s*\S")
ERROR_RE = re.compile(r"^#[A-Z0-9/_]+[!?]?$")
REWARD_LIKE_RE = re.compile(r"^x?\d", re.IGNORECASE)
# Ingredient results are often drawn as icons via =XLOOKUP("<name>", Images!...),
# which have no cached text value.
ICON_RE = re.compile(r'XLOOKUP\(\s*"([^"]+)"\s*,\s*Images!', re.IGNORECASE)
SHEET_ID_RE = re.compile(r"docs\.google\.com/spreadsheets/d/([\w-]+)")
GID_RE = re.compile(r"[?&#]gid=(\d+)")


@dataclass
class Cell:
    text: str
    ref: str = ""


@dataclass
class Sheet:
    name: str
    rows: list[list[Cell]]

    def at(self, row: int, col: int) -> Cell:
        if 0 <= row < len(self.rows) and 0 <= col < len(self.rows[row]):
            return self.rows[row][col]
        return Cell("")


@dataclass
class Matrix:
    sheet: str
    ingredients: list[str]
    cells: list[list[Cell]]
    location: str


@dataclass
class Audit:
    row_ingredient: str = ""
    column_ingredient: str = ""
    source_value: str = ""
    normalised_value: str = ""
    classification: str = "blank"
    quantity: int | None = None
    optimizer_key: str = ""
    mapping_rule: str = ""
    was_changed: bool = False
    was_mirrored: bool = False
    cell: str = ""
    parsed_reward: str = ""


@dataclass
class CleanResult:
    source: str
    matrix: str
    ingredients: list[str]
    audits: list[Audit]
    counts: Counter
    invalid: list[tuple[str, str, str]]
    conflicts: list[tuple[str, str, str, str]]
    problems: list[str]
    diff: dict
    dry_run: bool
    target: Path
    updated: bool = False
    valid_cells: int = field(init=False)

    def __post_init__(self):
        bad = {(r, c) for r, c, _kind in self.invalid}
        for r, c, _left, _right in self.conflicts:
            bad |= {(r, c), (c, r)}
        self.valid_cells = self.cell_count - len(bad)

    @property
    def cell_count(self) -> int:
        return len(self.ingredients) ** 2

    @property
    def passed(self) -> bool:
        return self.valid_cells == self.cell_count and not self.problems

    @property
    def mapping_coverage(self) -> float:
        return 100.0 * self.valid_cells / self.cell_count if self.cell_count else 0.0

    @property
    def normalisations(self) -> Counter:
        return Counter(
            a.mapping_rule
            for a in self.audits
            if a.was_changed and not a.was_mirrored and a.classification in VALID and a.mapping_rule != "identity"
        )

    @property
    def mirrored(self) -> list[Audit]:
        return [a for a in self.audits if a.was_mirrored]

    def issues(self) -> list[dict]:
        audits = {(a.row_ingredient, a.column_ingredient): a for a in self.audits}
        issues = []
        for row, col, kind in self.invalid:
            audit = audits.get((row, col), Audit())
            issues.append({
                "row_ingredient": row,
                "column_ingredient": col,
                "classification": kind,
                "source_value": clean_text(audit.source_value),
                "parsed_reward": audit.parsed_reward,
                "cell": audit.cell,
            })
        return issues


def load_reward_keys(streamlit_path: Path | str) -> list[str]:
    path = Path(streamlit_path)
    try:
        tree = ast.parse(path.read_text(encoding="utf-8"))
    except (OSError, SyntaxError) as exc:
        raise CleanerError(f"Could not parse {path} to read default_importance_scores: {exc}") from exc
    for node in tree.body:
        if isinstance(node, ast.Assign) and any(
            isinstance(t, ast.Name) and t.id == "default_importance_scores" for t in node.targets
        ):
            try:
                scores = ast.literal_eval(node.value)
            except ValueError:
                break
            if isinstance(scores, dict) and scores and all(isinstance(k, str) for k in scores):
                return list(scores)
            break
    raise CleanerError(
        f"Could not extract a literal default_importance_scores dict from {path.name}. "
        "Refusing to fall back to a hard-coded reward list."
    )


def clean_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_google_sheets_url(url: str) -> tuple[str, str]:
    sheet_id, gid = SHEET_ID_RE.search(url), GID_RE.search(url)
    if not sheet_id:
        raise CleanerError(f"Not a Google Sheets URL: {url}")
    if not gid:
        raise CleanerError("The Google Sheets URL has no gid. Copy the URL while the recipe tab is open.")
    return sheet_id.group(1), gid.group(1)


def download_google_sheet(url: str, timeout: int = 60) -> bytes:
    """Download only the gid tab as xlsx, which keeps the icon formulas the CSV export drops."""
    sheet_id, gid = parse_google_sheets_url(url)
    export = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx&gid={gid}"
    try:
        with urllib.request.urlopen(export, timeout=timeout) as response:
            payload = response.read()
    except urllib.error.HTTPError as exc:
        raise CleanerError(
            f"Google Sheets export failed (HTTP {exc.code}). Share the sheet as 'Anyone with the link can view'."
        ) from exc
    except urllib.error.URLError as exc:
        raise CleanerError(f"Could not reach Google Sheets: {exc.reason}") from exc
    if not payload.startswith(b"PK"):
        if not payload.strip():
            kind = "an empty response"
        elif b"<html" in payload[:4000].lower():
            kind = "an HTML login or error page"
        else:
            kind = "something that is not a spreadsheet"
        raise CleanerError(f"Google Sheets returned {kind} instead of the tab export. Check the sharing settings.")
    return payload


def load_source(source: str, sheet: str | None = None) -> list[Sheet]:
    if re.match(r"https?://", source):
        if sheet:
            raise CleanerError("--sheet only applies to local .xlsx files.")
        sheets = load_xlsx(io.BytesIO(download_google_sheet(source)))
        if len(sheets) != 1:
            raise CleanerError(f"Expected Google to export one tab, got {len(sheets)}.")
        return sheets
    path = Path(source)
    if not path.is_file():
        raise CleanerError(f"Source file not found: {source}")
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return load_xlsx(path, sheet)
    if suffix == ".csv" and not sheet:
        rows = csv.reader(io.StringIO(path.read_text(encoding="utf-8-sig")))
        return [Sheet(path.name, [
            [Cell(clean_text(v), f"{get_column_letter(c)}{r}") for c, v in enumerate(row, 1)]
            for r, row in enumerate(rows, 1)
        ])]
    raise CleanerError("Use a Google Sheets URL, a .csv file, or a .xlsx file (--sheet is xlsx only).")


def load_xlsx(file, sheet: str | None = None) -> list[Sheet]:
    try:
        values = openpyxl.load_workbook(file, data_only=True)
        if hasattr(file, "seek"):
            file.seek(0)
        formulas = openpyxl.load_workbook(file)
    except Exception as exc:
        raise CleanerError(f"Could not read workbook: {exc}") from exc
    names = values.sheetnames
    if sheet is not None:
        if sheet not in names:
            raise CleanerError(f"Worksheet {sheet!r} not found. Sheets: {', '.join(names)}")
        names = [sheet]
    sheets = []
    for name in names:
        rows = [
            [Cell(_cell_text(v.value, f.value), v.coordinate) for v, f in zip(value_row, formula_row)]
            for value_row, formula_row in zip(values[name].iter_rows(), formulas[name].iter_rows())
        ]
        sheets.append(Sheet(name, rows))
    return sheets


def _cell_text(value, formula) -> str:
    text = clean_text(value)
    if text and not ERROR_RE.match(text):
        return text
    icon = ICON_RE.search(str(getattr(formula, "text", formula) or ""))
    return icon.group(1) if icon else text


def _is_label(text: str) -> bool:
    return bool(text) and not ERROR_RE.match(text) and not REWARD_LIKE_RE.match(text)


def find_matrices(sheet: Sheet) -> list[Matrix]:
    """Each maximal run of header labels paired with a column holding the same names below it."""
    found = []
    for r, row in enumerate(sheet.rows):
        c = 0
        while c < len(row):
            if not _is_label(row[c].text):
                c += 1
                continue
            start = c
            while c < len(row) and _is_label(row[c].text):
                c += 1
            matrix = _match_row_labels(sheet, r, start, [cell.text for cell in row[start:c]])
            if matrix:
                found.append(matrix)
    return found


def _match_row_labels(sheet: Sheet, header_row: int, start: int, names: list[str]) -> Matrix | None:
    n = len(names)
    wanted = sorted(name.casefold() for name in names)
    if n < 2 or len(set(wanted)) != n:
        return None
    for col in range(start - 1, -1, -1):
        column = [sheet.at(r, col).text.casefold() for r in range(len(sheet.rows))]
        for top in range(header_row + 1, len(column) - n + 1):
            if sorted(column[top:top + n]) == wanted:
                row_of = {label: top + i for i, label in enumerate(column[top:top + n])}
                cells = [[sheet.at(row_of[a.casefold()], start + j) for j in range(n)] for a in names]
                return Matrix(sheet.name, names, cells, f"{sheet.name}!{sheet.rows[header_row][start].ref}")
    return None


def detect_matrix(sheets: list[Sheet]) -> Matrix:
    ranked = []
    for sheet in sheets:
        found = find_matrices(sheet)
        if not found:
            continue
        size = max(len(m.ingredients) for m in found)
        largest = [m for m in found if len(m.ingredients) == size]
        archived = any(term in sheet.name.casefold() for term in ARCHIVE_TERMS)
        for m in largest:
            filled = sum(bool(cell.text) for row in m.cells for cell in row)
            ranked.append(((not archived, size, len(largest) == 1, filled), m))
    if not ranked:
        raise CleanerError("No ingredient x ingredient recipe matrix found (header row plus matching row labels).")
    ranked.sort(key=lambda item: item[0], reverse=True)
    best = [m for score, m in ranked if score == ranked[0][0]]
    if len(best) > 1:
        raise CleanerError(
            "Several equally plausible recipe matrices found; pass --sheet to choose one:\n"
            + "\n".join(f"- {m.location}" for m in best)
        )
    return best[0]


def normalise_recipe_value(source, ingredients: list[str], keys: list[str]) -> Audit:
    raw = "" if source is None else str(source)
    text = clean_text(raw)

    def audit(classification, value=text, rule="identity", quantity=None, key="", parsed=""):
        return Audit(
            source_value=raw, normalised_value=value, classification=classification, quantity=quantity,
            optimizer_key=key, mapping_rule=rule, was_changed=value != raw, parsed_reward=parsed,
        )

    names = {name.casefold(): name for name in ingredients}
    if not text:
        return audit("blank", "", "blank")
    if text.casefold() in EXACT_VALUES:
        override = normalise_recipe_value(EXACT_VALUES[text.casefold()], ingredients, keys)
        override.source_value, override.was_changed = raw, override.normalised_value != raw
        override.mapping_rule = f"{text} -> {override.normalised_value}"
        return override
    if COMPOUND_RE.search(text):
        return audit("compound_reward", rule="compound_reward", parsed=text)
    if text.casefold() in names:
        name = names[text.casefold()]
        return audit("ingredient", name, "identity" if name == text else f"{text} -> {name}")

    match = QUANTITY_RE.match(text)
    quantity, description = (int(match[1]), match[2]) if match else (1, text)
    folded = description.casefold()
    if match and folded in names:
        return audit("ambiguous", rule="quantity_on_ingredient", parsed=description)

    key, rule = None, ""
    if folded in {k.casefold(): k for k in keys}:
        key = {k.casefold(): k for k in keys}[folded]
    elif folded in REWARD_ALIASES:
        key = REWARD_ALIASES[folded]
        rule = f"{description} -> {key}"
    else:
        for pattern, target in EQUIPMENT_VARIANTS:
            if pattern.match(description):
                key = target
                rule = f"{description.split('(')[0].strip()} (...) -> {target}"
                break
    if key not in keys or quantity < 1:
        base = re.sub(r"\s*\(.*\)$", "", description) or description
        return audit("unsupported", rule="unmapped", quantity=quantity if match else None, parsed=base)

    value = f"{quantity} {key}"
    if not rule:
        rule = "identity" if value == text else f"{text} -> {value}"
    return audit("supported_reward", value, rule, quantity, key, key)


def clean_matrix(matrix: Matrix, keys: list[str]) -> list[list[Audit]]:
    grid = []
    for r, row_name in enumerate(matrix.ingredients):
        row = []
        for c, col_name in enumerate(matrix.ingredients):
            cell = matrix.cells[r][c]
            audit = normalise_recipe_value(cell.text, matrix.ingredients, keys)
            audit.row_ingredient, audit.column_ingredient = row_name, col_name
            audit.cell = f"{matrix.sheet}!{cell.ref}" if cell.ref else ""
            row.append(audit)
        grid.append(row)
    for i in range(len(grid)):
        for j in range(i + 1, len(grid)):
            a, b = grid[i][j], grid[j][i]
            if (a.classification == "blank") != (b.classification == "blank"):
                src, dst = (b, a) if a.classification == "blank" else (a, b)
                for name in ("normalised_value", "classification", "quantity", "optimizer_key", "parsed_reward"):
                    setattr(dst, name, getattr(src, name))
                dst.mapping_rule, dst.was_mirrored, dst.was_changed = "mirror", True, True
    return grid


def classify_canonical(value, ingredients: list[str], keys: list[str]) -> tuple[str, int | None, str | None]:
    """Strict final check: exact ingredient or '<positive int> <exact key>'. No aliases."""
    text = "" if value is None else str(value)
    if not text.strip():
        return "blank", None, None
    if text in ingredients:
        return "ingredient", None, None
    if COMPOUND_RE.search(text):
        return "compound_reward", None, None
    match = CANONICAL_RE.match(text)
    if match and match[2] in keys:
        return "supported_reward", int(match[1]), match[2]
    return "unsupported", None, None


def validate_frame(frame: pd.DataFrame, keys: list[str]):
    ingredients = list(frame.index)
    problems = []
    if list(frame.columns) != ingredients:
        problems.append("Candidate columns do not match its row ingredients.")
    if len(set(ingredients)) != len(ingredients) or not all(ingredients):
        problems.append("Candidate ingredient names are blank or duplicated.")
    counts, invalid, conflicts = Counter(), [], []
    for row in ingredients:
        for col in frame.columns:
            kind = classify_canonical(frame.at[row, col], ingredients, keys)[0]
            counts[kind] += 1
            if kind not in VALID:
                invalid.append((row, col, kind))
    if not problems:
        for i, row in enumerate(ingredients):
            for col in ingredients[i + 1:]:
                left, right = frame.at[row, col], frame.at[col, row]
                if left and right and left != right:
                    conflicts.append((row, col, left, right))
    return counts, invalid, conflicts, problems


def read_recipe_csv(path: Path) -> pd.DataFrame:
    try:
        frame = pd.read_csv(path, index_col=0, dtype=str, keep_default_na=False)
    except Exception as exc:
        raise CleanerError(f"Could not read {path}: {exc}") from exc
    frame.index = frame.index.astype(str)
    frame.index.name = None
    return frame


def reward_types(frame: pd.DataFrame, keys: list[str]) -> list[str]:
    ingredients = list(frame.index)
    found = {normalise_recipe_value(v, ingredients, keys).optimizer_key for v in frame.to_numpy().ravel()}
    return [k for k in keys if k in found]


def compare_frames(old: pd.DataFrame | None, new: pd.DataFrame, keys: list[str]) -> dict:
    """Pairwise comparison after normalising the old CSV, so tidy-ups are not counted as changes."""
    new_types = reward_types(new, keys)
    if old is None:
        return {"had_existing": False, "new_ingredient_count": len(new.index), "new_reward_types": new_types}
    old_names, new_names = list(old.index), list(new.index)
    shared = [name for name in new_names if name in old_names and name in old.columns]
    unchanged = sum(
        normalise_recipe_value(old.at[r, c], old_names, keys).normalised_value == new.at[r, c]
        for r in shared
        for c in shared
    )
    old_types = reward_types(old, keys)
    return {
        "had_existing": True,
        "old_ingredient_count": len(old_names),
        "new_ingredient_count": len(new_names),
        "ingredients_added": [n for n in new_names if n not in old_names],
        "ingredients_removed": [n for n in old_names if n not in new_names],
        "cells_unchanged": unchanged,
        "cells_changed": len(new_names) ** 2 - unchanged,
        "old_reward_types": old_types,
        "new_reward_types": new_types,
        "reward_types_added": [k for k in new_types if k not in old_types],
        "reward_types_removed": [k for k in old_types if k not in new_types],
    }


def atomic_replace(data: bytes, target: Path) -> None:
    fd, tmp = tempfile.mkstemp(dir=target.parent, prefix=".alchemy-", suffix=".csv")
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(data)
        os.replace(tmp, target)
    except BaseException:
        Path(tmp).unlink(missing_ok=True)
        raise


def write_reports(result: CleanResult, report_dir: Path) -> None:
    with (report_dir / "mapping_audit.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(asdict(Audit())))
        writer.writeheader()
        writer.writerows(asdict(a) for a in result.audits)
    report = {
        "passed": result.passed,
        "dry_run": result.dry_run,
        "updated": result.updated,
        "source": result.source,
        "matrix": result.matrix,
        "ingredients": result.ingredients,
        "cell_count": result.cell_count,
        "valid_cells": result.valid_cells,
        "mapping_coverage": result.mapping_coverage,
        "counts": dict(result.counts),
        "normalisations": dict(result.normalisations.most_common()),
        "mirrored": [f"{a.row_ingredient} + {a.column_ingredient}" for a in result.mirrored],
        "issues": result.issues(),
        "symmetry_conflicts": [
            {"pair": f"{r} + {c}", "left": left, "right": right} for r, c, left, right in result.conflicts
        ],
        "problems": result.problems,
        "diff": result.diff,
    }
    (report_dir / "validation_report.json").write_text(json.dumps(report, indent=2), encoding="utf-8")


def run_clean(
    source: str,
    *,
    target: Path,
    report_dir: Path,
    streamlit_path: Path,
    dry_run: bool = False,
    sheet: str | None = None,
) -> CleanResult:
    """Clean ``source`` and replace ``target`` only if every cell validates."""
    keys = load_reward_keys(streamlit_path)
    matrix = detect_matrix(load_source(source, sheet))
    grid = clean_matrix(matrix, keys)

    report_dir.mkdir(parents=True, exist_ok=True)
    candidate_path = report_dir / "candidate.csv"
    values = [[a.normalised_value for a in row] for row in grid]
    pd.DataFrame(values, index=matrix.ingredients, columns=matrix.ingredients).to_csv(
        candidate_path, encoding="utf-8", lineterminator="\n"
    )
    candidate = read_recipe_csv(candidate_path)
    counts, invalid, conflicts, problems = validate_frame(candidate, keys)
    old = read_recipe_csv(target) if target.is_file() else None

    result = CleanResult(
        source=source,
        matrix=matrix.location,
        ingredients=list(matrix.ingredients),
        audits=[a for row in grid for a in row],
        counts=counts,
        invalid=invalid,
        conflicts=conflicts,
        problems=problems,
        diff=compare_frames(old, candidate, keys),
        dry_run=dry_run,
        target=target,
    )
    result.updated = result.passed and not dry_run
    write_reports(result, report_dir)
    if result.updated:
        atomic_replace(candidate_path.read_bytes(), target)
    return result
