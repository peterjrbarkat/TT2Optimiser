"""Tests for the deterministic Alchemy recipe cleaner (python -m unittest discover -s tests -v)."""

from __future__ import annotations

import io
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import openpyxl
import pandas as pd

from src.alchemy_cleaner import (
    CleanerError,
    classify_canonical,
    load_reward_keys,
    normalise_recipe_value,
    parse_google_sheets_url,
    run_clean,
)
from update_alchemy_csv import render_summary

ROOT = Path(__file__).resolve().parents[1]
STREAMLIT = ROOT / "streamlit_app.py"
WORKBOOK = ROOT / "TT2 Alchemy Event_20_06_26.xlsx"
PRODUCTION = ROOT / "TT2 Alchemy Event.csv"
SHEET_URL = (
    "https://docs.google.com/spreadsheets/d/"
    "1ZIn5pv6_TBeUNr1ze9_kjE4uKWicmVWQYdsMaBOAQNo/edit?gid=1745300311#gid=1745300311"
)
KEYS = load_reward_keys(STREAMLIT)


def extract_loot(value, keys):
    """Copy of streamlit_app.extract_loot, used to check optimiser equivalence."""
    if isinstance(value, str):
        parts = value.strip().split()
        try:
            amount = int(parts[0])
            item_type = " ".join(parts[1:])
            return next(((k, amount) for k in keys if k in item_type), (item_type, amount))
        except (ValueError, IndexError):
            for key in sorted(keys, key=len, reverse=True):
                if key in value:
                    return (key, 1)
    return ("Unknown", 0)


def normalise(value, ingredients=("Crystal", "Leaf")):
    return normalise_recipe_value(value, list(ingredients), KEYS)


def fake_response(payload: bytes):
    response = mock.MagicMock()
    response.__enter__.return_value.read.return_value = payload
    return response


class TempRun(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)
        self.target = self.root / "TT2 Alchemy Event.csv"
        shutil.copyfile(PRODUCTION, self.target)
        self.before = self.target.read_bytes()

    def run_source(self, source, **kwargs):
        return run_clean(
            str(source), target=self.target, report_dir=self.root / ".alchemy", streamlit_path=STREAMLIT, **kwargs
        )

    def run_csv(self, text, **kwargs):
        path = self.root / "source.csv"
        path.write_text(text, encoding="utf-8")
        return self.run_source(path, **kwargs)


class RewardKeyTests(unittest.TestCase):
    def test_keys_come_from_streamlit_app(self):
        self.assertIn("Eggs", KEYS)
        self.assertIn("Clan Scroll", KEYS)
        self.assertNotIn("Dust", KEYS)

    def test_missing_dictionary_fails_clearly(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "streamlit_app.py"
            path.write_text("scores = {'Eggs': 1}\n", encoding="utf-8")
            with self.assertRaisesRegex(CleanerError, "default_importance_scores"):
                load_reward_keys(path)

    def test_cleaner_never_imports_the_app(self):
        source = (ROOT / "src" / "alchemy_cleaner.py").read_text(encoding="utf-8")
        self.assertNotRegex(source, r"(?m)^\s*(import|from)\s+streamlit_app\b")


class NormalisationTests(unittest.TestCase):
    def assertMaps(self, source, expected):
        result = normalise(source)
        self.assertIn(result.classification, ("ingredient", "supported_reward"), source)
        self.assertEqual(result.normalised_value, expected, source)

    def test_ingredient_case(self):
        self.assertMaps("crystal", "Crystal")

    def test_reward_aliases(self):
        for source, expected in {
            "8 Pet Eggs": "8 Eggs",
            "x8 Pet Eggs": "8 Eggs",
            "115 Wild Cards": "115 Wildcards",
            "1 Skill Point": "1 Skill Points",
            "3 Clan Scrolls": "3 Clan Scroll",
            "2 Fortune Scrolls": "2 Fortune Scroll",
            "13 Fortune Hero Weapons": "13 Hero Weapons",
        }.items():
            self.assertMaps(source, expected)

    def test_equipment_variants_and_typo(self):
        for source in ("Event Equipment (Lab Coat)", "Safety Tongs (Event Set)"):
            self.assertMaps(source, "1 Event Equipment")
        for source in ("Legendary Equipment (Madcap Goggles)", "Legendary Equip (Shadowgold Armor)", "Legendary Insult"):
            self.assertMaps(source, "1 Legendary Equipment")

    def test_unsupported_terms_are_not_guessed(self):
        for source, parsed in {
            "1 Perk": "Perk",
            "3 Perks": "Perks",
            "1020 Dust": "Dust",
            "Mythic Equipment (Divine Cleave)": "Mythic Equipment",
            "Unique Equipment (Maple Brush)": "Unique Equipment",
            "4 Moon Rocks": "Moon Rocks",
        }.items():
            result = normalise(source)
            self.assertEqual(result.classification, "unsupported", source)
            self.assertEqual(result.parsed_reward, parsed)
            self.assertEqual(result.optimizer_key, "")

    def test_compound_reward_is_rejected(self):
        result = normalise("Event Avatar + 97 Currency")
        self.assertEqual(result.classification, "compound_reward")
        self.assertNotEqual(result.normalised_value, "97 Currency")

    def test_approved_exact_values(self):
        self.assertMaps("Avatar + 108 Currency", "108 Currency")
        self.assertMaps("Unique Equipment (The Retaliator Reforged)", "1 Event Equipment")
        self.assertEqual(normalise("Avatar + 110 Currency").classification, "compound_reward")

    def test_final_check_ignores_aliases(self):
        self.assertEqual(classify_canonical("8 Pet Eggs", ["Leaf"], KEYS)[0], "unsupported")
        self.assertEqual(classify_canonical("8 Eggs", ["Leaf"], KEYS), ("supported_reward", 8, "Eggs"))


class PipelineTests(TempRun):
    def test_one_sided_blank_is_mirrored_and_written(self):
        result = self.run_csv(",A,B\nA,1 Currency,1 Currency\nB,,2 Eggs\n")
        self.assertTrue(result.passed, render_summary(result))
        self.assertEqual([(a.row_ingredient, a.column_ingredient) for a in result.mirrored], [("B", "A")])
        self.assertEqual(pd.read_csv(self.target, index_col=0).loc["B", "A"], "1 Currency")

    def test_both_sides_blank_fails(self):
        result = self.run_csv(",A,B\nA,1 Currency,\nB,,2 Eggs\n")
        self.assertFalse(result.passed)
        self.assertEqual(result.counts["blank"], 2)
        self.assertEqual(self.target.read_bytes(), self.before)

    def test_symmetry_conflict_fails(self):
        result = self.run_csv(",A,B\nA,1 Currency,2 Currency\nB,3 Eggs,1 Currency\n")
        self.assertFalse(result.passed)
        self.assertEqual(len(result.conflicts), 1)
        self.assertEqual(self.target.read_bytes(), self.before)

    def test_failed_validation_leaves_production_bytes_unchanged(self):
        result = self.run_csv(",A,B\nA,1 Perk,Event Avatar + 97 Currency\nB,Event Avatar + 97 Currency,2 Eggs\n")
        self.assertFalse(result.passed)
        self.assertFalse(result.updated)
        self.assertEqual(self.target.read_bytes(), self.before)
        for name in ("candidate.csv", "mapping_audit.csv", "validation_report.json"):
            self.assertTrue((self.root / ".alchemy" / name).is_file())
        self.assertIn("NOT UPDATED", render_summary(result))

    def test_dry_run_never_writes(self):
        result = self.run_csv(",A,B\nA,B,1 Currency\nB,1 Currency,2 Eggs\n", dry_run=True)
        self.assertTrue(result.passed)
        self.assertEqual(self.target.read_bytes(), self.before)
        self.assertIn("MAPPING COVERAGE: 100%", render_summary(result))

    def test_row_order_follows_header(self):
        self.run_csv(",Blue,Red\nRed,1 Currency,Red\nBlue,Blue,1 Currency\n")
        frame = pd.read_csv(self.target, index_col=0)
        self.assertEqual(list(frame.index), ["Blue", "Red"])
        self.assertEqual(frame.loc["Blue", "Blue"], "Blue")

    def test_archive_sheet_loses_and_ties_fail(self):
        path = self.root / "book.xlsx"
        workbook = openpyxl.Workbook()
        for index, title in enumerate(("Current", "Old copy")):
            sheet = workbook.active if index == 0 else workbook.create_sheet()
            sheet.title = title
            sheet["A1"] = "Recipe book title"
            sheet["C3"], sheet["D3"] = "Red", "Blue"
            sheet["A5"], sheet["A6"] = "Red", "Blue"
            sheet["C5"], sheet["D5"], sheet["C6"], sheet["D6"] = "Blue", "1 Currency", "1 Currency", "3 Eggs"
        workbook.save(path)
        result = self.run_source(path, dry_run=True)
        self.assertTrue(result.passed, render_summary(result))
        self.assertTrue(result.matrix.startswith("Current!"))

        workbook["Old copy"].title = "Also current"
        workbook.save(path)
        with self.assertRaisesRegex(CleanerError, "equally plausible"):
            self.run_source(path, dry_run=True)


class GoogleSheetTests(TempRun):
    def icon_workbook(self) -> bytes:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "8.2 Recipe Matrix"
        sheet["C3"], sheet["D3"] = "Red", "Blue"
        sheet["A5"], sheet["A6"] = "Red", "Blue"
        sheet["C4"] = '=_xlfn.XLOOKUP(C3,Images!$A:$A,Images!$C:$C,"")'
        sheet["C5"] = '=_xlfn.XLOOKUP("Blue",Images!$A:$A,Images!$C:$C,"")'
        sheet["D5"] = "x3 Pet Eggs"
        sheet["C6"] = "=INDEX($C$5:$D$6, COLUMN()-2, ROW()-4)"
        sheet["D6"] = "2 Fortune Scrolls"
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def test_url_keeps_requested_gid(self):
        self.assertEqual(
            parse_google_sheets_url(SHEET_URL), ("1ZIn5pv6_TBeUNr1ze9_kjE4uKWicmVWQYdsMaBOAQNo", "1745300311")
        )
        with self.assertRaisesRegex(CleanerError, "gid"):
            parse_google_sheets_url("https://docs.google.com/spreadsheets/d/abc/edit")

    def test_url_update_resolves_icon_formulas_and_replaces_csv(self):
        with mock.patch("urllib.request.urlopen", return_value=fake_response(self.icon_workbook())) as urlopen:
            result = self.run_source(SHEET_URL)
        self.assertIn("export?format=xlsx&gid=1745300311", urlopen.call_args[0][0])
        self.assertTrue(result.passed, render_summary(result))
        self.assertTrue(result.updated)
        frame = pd.read_csv(self.target, index_col=0)
        self.assertEqual(frame.loc["Red", "Red"], "Blue")
        self.assertEqual(frame.loc["Blue", "Red"], "3 Eggs")
        self.assertEqual(frame.loc["Blue", "Blue"], "2 Fortune Scroll")

    def test_html_or_empty_response_fails_without_writing(self):
        for payload in (b"<!DOCTYPE html><html>Sign in</html>", b""):
            with mock.patch("urllib.request.urlopen", return_value=fake_response(payload)):
                with self.assertRaisesRegex(CleanerError, "instead of the tab export"):
                    self.run_source(SHEET_URL)
            self.assertEqual(self.target.read_bytes(), self.before)


class JuneWorkbookRegression(TempRun):
    def test_june_workbook_is_fully_mapped_and_equivalent(self):
        result = self.run_source(WORKBOOK, dry_run=True)
        summary = render_summary(result)
        self.assertTrue(result.passed, summary)
        self.assertIn("MAPPING COVERAGE: 100%", summary)
        self.assertTrue(result.matrix.startswith("Recipe Matrix!"))
        self.assertEqual((result.cell_count, result.valid_cells, result.counts["unsupported"]), (256, 256, 0))
        self.assertEqual(self.target.read_bytes(), self.before)

        # The workbook's Output sheet holds the CSV published for the June event.
        output = openpyxl.load_workbook(WORKBOOK, data_only=True)["Output"]
        published = "\n".join(str(row[0]) for row in output.iter_rows(min_col=2, max_col=2, values_only=True) if row[0])
        old = pd.read_csv(io.StringIO(published), index_col=0, dtype=str)
        new = pd.read_csv(self.root / ".alchemy" / "candidate.csv", index_col=0, dtype=str)
        self.assertEqual(list(new.index), list(old.index))
        self.assertEqual(list(new.columns), list(old.index))
        ingredients = list(new.index)
        for row in ingredients:
            for col in ingredients:
                old_value, new_value = old.at[row, col], new.at[row, col]
                kind, _qty, _key = classify_canonical(new_value, ingredients, KEYS)
                self.assertIn(kind, ("ingredient", "supported_reward"), f"{row} + {col}")
                if kind == "ingredient":
                    self.assertEqual(new_value.casefold(), old_value.casefold(), f"{row} + {col}")
                else:
                    self.assertEqual(extract_loot(new_value, KEYS), extract_loot(old_value, KEYS), f"{row} + {col}")
        self.assertEqual(result.normalisations["crystal -> Crystal"], 2)


if __name__ == "__main__":
    unittest.main()
